"""Create a test .xlsx workbook for the excel translation command.

Builds a workbook with:
- Arabic legal text in cells (translatable)
- English legal text in cells (translatable)
- A formula (must be preserved, NOT translated)
- A merged cell with text
- A cell comment
- A header row
- A chart title
- Numbers (must be preserved, not translated)
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.chart import BarChart, Reference


def main() -> None:
    wb: Workbook = Workbook()

    # --- Sheet 1: Arabic legal articles (ar-en direction) ---
    ws = wb.active
    ws.title = "Civil_Code_AR"

    # Header row
    ws["A1"] = "رقم المادة"
    ws["B1"] = "نص المادة"
    ws["C1"] = "الملاحظات"

    # Article 1
    ws["A2"] = 1
    ws["B2"] = "عقد البيع هو اتفاق يلتزم بمقتضاه البائع أن ينقل للمشتري ملكية شيء أو حق مالي آخر مقابل ثمن نقدي."
    ws["C2"].comment = Comment("هذه المادة تعرّف عقد البيع", "LegalTranslator")

    # Article 2
    ws["A3"] = 2
    ws["B3"] = "يتم العقد بمجرد تبادل الإيجاب والقبول بين الطرفين."
    ws["C3"].comment = Comment("انظر المادة 148 من القانون المدني", "LegalTranslator")

    # Article 3 — merged cell
    ws.merge_cells("B4:C4")
    ws["A4"] = 3
    ws["B4"] = "إذا كان الشيء المبيع معيناً بالذات وجب تسليمه في الحال ما لم يوجد اتفاق على غير ذلك."

    # Article 4 — formula (must be preserved, NOT translated)
    ws["A5"] = 4
    ws["B5"] = "يلتزم البائع بضمان التعرض والاستحقاق."
    ws["C5"] = "=A2+A3+A4"  # formula — must survive translation intact

    # --- Sheet 2: English legal articles (en-ar direction) ---
    ws2 = wb.create_sheet("Commercial_Code_EN")

    ws2["A1"] = "Article No."
    ws2["B1"] = "Article Text"
    ws2["C1"] = "Notes"

    ws2["A2"] = 1
    ws2["B2"] = "A contract of sale is an agreement whereby the seller undertakes to transfer to the buyer the ownership of a thing or another financial right in consideration of a cash price."
    ws2["C2"].comment = Comment("This article defines a contract of sale", "LegalTranslator")

    ws2["A3"] = 2
    ws2["B3"] = "The contract is concluded upon the exchange of offer and acceptance between the two parties."
    ws2["C3"].comment = Comment("See Article 148 of the Civil Code", "LegalTranslator")

    ws2["A4"] = 3
    ws2["B4"] = "If the sold item is specifically identified, it must be delivered immediately unless otherwise agreed."
    ws2["C4"] = "=A2+A3"  # formula — must survive translation intact

    # --- Sheet 3: Chart (title should be translated) ---
    ws3 = wb.create_sheet("Statistics")
    ws3["A1"] = "السنة"
    ws3["B1"] = "عدد القضايا"
    ws3["A2"] = 2020
    ws3["B2"] = 150
    ws3["A3"] = 2021
    ws3["B3"] = 200
    ws3["A4"] = 2022
    ws3["B4"] = 180

    chart = BarChart()
    chart.title = "إحصائيات القضايا"
    chart.x_axis.title = "السنة"
    chart.y_axis.title = "العدد"
    data = Reference(ws3, min_col=2, min_row=1, max_row=4, max_col=2)
    cats = Reference(ws3, min_col=1, min_row=2, max_row=4)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    ws3.add_chart(chart, "D2")

    out: Path = Path("data/test_legal_workbook.xlsx")
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out))
    print(f"Created: {out} ({out.stat().st_size} bytes)")
    print("Sheets: Civil_Code_AR, Commercial_Code_EN, Statistics")
    print("Features: Arabic text, English text, formulas, merged cells, comments, chart")


if __name__ == "__main__":
    main()
